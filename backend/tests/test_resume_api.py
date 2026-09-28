from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from tests.conftest import auth_header, register_user

SAMPLE = (
    Path(__file__).resolve().parents[2] / "datasets" / "synthetic" / "sample_resume.txt"
)


def test_resume_upload_extracts_inferred_skills(
    client: TestClient, unique_email: str
) -> None:
    created = register_user(client, unique_email)
    headers = auth_header(created["tokens"]["access_token"])
    files = {
        "file": ("sample_resume.txt", SAMPLE.read_bytes(), "text/plain"),
    }
    response = client.post("/api/v1/employees/me/resumes", headers=headers, files=files)
    assert response.status_code == 201, response.text
    body = response.json()["data"]
    names = {item["skill"]["canonical_name"] for item in body["skills"]}
    assert "Machine Learning" in names
    assert "Kubernetes" in names
    assert "Natural Language Processing" in names
    assert all(item["inferred"] is True for item in body["skills"])
    assert all(item["source_type"] == "RESUME" for item in body["skills"])
    assert all(item["evidence_snippet"] for item in body["skills"])

    listed = client.get("/api/v1/employees/me/resumes", headers=headers)
    assert listed.status_code == 200
    assert listed.json()["data"][0]["skill_count"] >= 1


def test_resume_rejects_unsupported_type(client: TestClient, unique_email: str) -> None:
    created = register_user(client, unique_email)
    headers = auth_header(created["tokens"]["access_token"])
    response = client.post(
        "/api/v1/employees/me/resumes",
        headers=headers,
        files={"file": ("photo.png", b"not-a-resume", "image/png")},
    )
    assert response.status_code == 422


def test_resume_docx_upload(client: TestClient, unique_email: str) -> None:
    created = register_user(client, unique_email)
    headers = auth_header(created["tokens"]["access_token"])
    document = Document()
    document.add_heading("Skills", level=1)
    document.add_paragraph("Python, Docker, ML")
    buffer = BytesIO()
    document.save(buffer)
    response = client.post(
        "/api/v1/employees/me/resumes",
        headers=headers,
        files={
            "file": (
                "resume.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 201, response.text
    names = {
        item["skill"]["canonical_name"] for item in response.json()["data"]["skills"]
    }
    assert "Python" in names
    assert "Machine Learning" in names


def test_resume_requires_auth(client: TestClient) -> None:
    response = client.post(
        "/api/v1/employees/me/resumes",
        files={"file": ("resume.txt", b"Skills\nPython\n", "text/plain")},
    )
    assert response.status_code == 401
