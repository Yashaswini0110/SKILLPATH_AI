"""Extract plain text from PDF, DOCX, and TXT resumes."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}
SUPPORTED_CONTENT_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "text/plain",
}


class UnsupportedResumeTypeError(ValueError):
    pass


def extract_text(path: str | Path, content_type: str | None = None) -> str:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise UnsupportedResumeTypeError(
            f"Unsupported resume type '{suffix or content_type}'. Use PDF, DOCX, or TXT."
        )
    if suffix == ".pdf":
        return _from_pdf(file_path)
    if suffix == ".docx":
        return _from_docx(file_path)
    return _from_txt(file_path)


def extract_text_from_bytes(data: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise UnsupportedResumeTypeError(
            f"Unsupported resume type '{suffix}'. Use PDF, DOCX, or TXT."
        )
    from tempfile import NamedTemporaryFile

    with NamedTemporaryFile(suffix=suffix, delete=False) as handle:
        handle.write(data)
        temp_path = Path(handle.name)
    try:
        return extract_text(temp_path)
    finally:
        temp_path.unlink(missing_ok=True)


def _from_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return _clean_extracted("\n".join(pages))


def _from_docx(path: Path) -> str:
    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" ".join(cell.text for cell in row.cells))
    return _clean_extracted("\n".join(parts))


def _from_txt(path: Path) -> str:
    return _clean_extracted(path.read_text(encoding="utf-8", errors="replace"))


def _clean_extracted(text: str) -> str:
    lines = [line.strip() for line in text.replace("\x00", " ").splitlines()]
    collapsed: list[str] = []
    for line in lines:
        if line or (collapsed and collapsed[-1] != ""):
            collapsed.append(line)
    return "\n".join(collapsed).strip()
